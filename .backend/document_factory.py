"""
Document Factory - Arabic PDF Generation System
Based on 06_DOCUMENT_FACTORY.md specification

Uses docxtpl for Word templating and LibreOffice/alternative for PDF conversion.
Supports Arabic RTL text with proper ligatures.
"""

import os
import subprocess
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Dict, Any, Optional, List
from dataclasses import dataclass

try:
    from docxtpl import DocxTemplate
    DOCXTPL_AVAILABLE = True
except ImportError:
    DOCXTPL_AVAILABLE = False
    print("Warning: docxtpl not installed. Run: pip install docxtpl")


@dataclass
class ReportData:
    """Data structure for report generation"""
    # Case information
    case_number: str
    court_name: str
    wilaya: str
    
    # Parties
    plaintiff_name: str
    defendant_name: str
    
    # Dates
    gregorian_date: str
    hijri_date: str
    assignment_date: str
    
    # Financial data
    basic_salary: Decimal
    gross_salary: Decimal
    eosb: Decimal
    notice_pay: Decimal
    leave_pay: Decimal
    total_due: Decimal
    
    # Service period
    years_of_service: int
    months_of_service: int
    
    # Additional
    expert_name: str = "طارق الخبير"
    expert_title: str = "خبير محاسب قانوني"


class DocumentFactory:
    """
    Factory for generating Arabic legal documents.
    
    Workflow:
    1. Load .docx template
    2. Fill placeholders using docxtpl
    3. Save filled document
    4. Convert to PDF (if LibreOffice available)
    """
    
    TEMPLATES_DIR = Path(__file__).parent / "templates"
    OUTPUT_DIR = Path(__file__).parent.parent / "public" / "generated_reports"
    
    def __init__(self):
        # Create directories if they don't exist
        self.TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        self.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        
        # Check LibreOffice availability
        self.libreoffice_path = self._find_libreoffice()
    
    def _find_libreoffice(self) -> Optional[str]:
        """Find LibreOffice executable for PDF conversion"""
        possible_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/usr/bin/soffice",
            "/usr/bin/libreoffice",
        ]
        
        for path in possible_paths:
            if os.path.exists(path):
                return path
        
        # Try system PATH
        try:
            result = subprocess.run(
                ["soffice", "--version"],
                capture_output=True,
                text=True,
                timeout=5
            )
            if result.returncode == 0:
                return "soffice"
        except (subprocess.TimeoutExpired, FileNotFoundError):
            pass
        
        return None
    
    def list_templates(self) -> List[str]:
        """List available templates"""
        if not self.TEMPLATES_DIR.exists():
            return []
        return [f.name for f in self.TEMPLATES_DIR.glob("*.docx")]
    
    def generate_report(
        self, 
        template_name: str, 
        data: ReportData,
        output_filename: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Generate a filled document from template.
        
        Args:
            template_name: Name of the template file (e.g., "report_template.docx")
            data: ReportData object with all fields
            output_filename: Optional custom output filename
            
        Returns:
            Dict with paths to generated files and status
        """
        if not DOCXTPL_AVAILABLE:
            return {
                "success": False,
                "error": "docxtpl library not installed"
            }
        
        template_path = self.TEMPLATES_DIR / template_name
        
        if not template_path.exists():
            return {
                "success": False,
                "error": f"Template not found: {template_name}"
            }
        
        # Generate output filename (sanitize case_number)
        if not output_filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_case_number = data.case_number.replace("/", "-").replace("\\", "-")
            output_filename = f"report_{safe_case_number}_{timestamp}"
        
        docx_output = self.OUTPUT_DIR / f"{output_filename}.docx"
        pdf_output = self.OUTPUT_DIR / f"{output_filename}.pdf"
        
        try:
            # Load template
            doc = DocxTemplate(template_path)
            
            # Prepare context
            context = self._data_to_context(data)
            
            # Render document
            doc.render(context)
            
            # Save filled document
            doc.save(docx_output)
            
            result = {
                "success": True,
                "docx_path": str(docx_output),
                "pdf_path": None,
                "message": "Document generated successfully"
            }
            
            # Convert to PDF if LibreOffice available
            if self.libreoffice_path:
                pdf_result = self._convert_to_pdf(docx_output, pdf_output)
                if pdf_result:
                    result["pdf_path"] = str(pdf_output)
                    result["message"] += " (PDF created)"
            else:
                result["warning"] = "LibreOffice not found - PDF conversion skipped"
            
            return result
            
        except Exception as e:
            return {
                "success": False,
                "error": str(e)
            }
    
    def _data_to_context(self, data: ReportData) -> Dict[str, Any]:
        """Convert ReportData to template context dictionary"""
        return {
            # Case
            "case_number": data.case_number,
            "court_name": data.court_name,
            "wilaya": data.wilaya,
            
            # Parties
            "plaintiff_name": data.plaintiff_name,
            "defendant_name": data.defendant_name,
            
            # Dates
            "gregorian_date": data.gregorian_date,
            "hijri_date": data.hijri_date,
            "assignment_date": data.assignment_date,
            
            # Financial - formatted with 3 decimal places
            "basic_salary": f"{data.basic_salary:.3f}",
            "gross_salary": f"{data.gross_salary:.3f}",
            "eosb": f"{data.eosb:.3f}",
            "notice_pay": f"{data.notice_pay:.3f}",
            "leave_pay": f"{data.leave_pay:.3f}",
            "total_due": f"{data.total_due:.3f}",
            
            # Service period
            "years_of_service": data.years_of_service,
            "months_of_service": data.months_of_service,
            "service_period": f"{data.years_of_service} سنة و {data.months_of_service} شهر",
            
            # Expert
            "expert_name": data.expert_name,
            "expert_title": data.expert_title,
            
            # Salary breakdown table (for {% for %} loops)
            "salary_table": [
                {"item": "الراتب الأساسي", "amount": f"{data.basic_salary:.3f}"},
                {"item": "إجمالي الراتب", "amount": f"{data.gross_salary:.3f}"},
            ],
            
            # EOSB breakdown table
            "eosb_table": [
                {"item": "مكافأة نهاية الخدمة", "amount": f"{data.eosb:.3f}"},
                {"item": "بدل الإشعار", "amount": f"{data.notice_pay:.3f}"},
                {"item": "رصيد الإجازات", "amount": f"{data.leave_pay:.3f}"},
                {"item": "المجموع", "amount": f"{data.total_due:.3f}"},
            ]
        }
    
    def _convert_to_pdf(self, docx_path: Path, pdf_path: Path) -> bool:
        """
        Convert DOCX to PDF using LibreOffice.
        
        LibreOffice is critical for Arabic RTL support.
        """
        if not self.libreoffice_path:
            return False
        
        try:
            # LibreOffice headless conversion
            cmd = [
                self.libreoffice_path,
                "--headless",
                "--convert-to", "pdf",
                "--outdir", str(pdf_path.parent),
                str(docx_path)
            ]
            
            result = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=60
            )
            
            # LibreOffice names output file same as input but with .pdf
            expected_pdf = pdf_path.parent / f"{docx_path.stem}.pdf"
            
            if expected_pdf.exists():
                # Rename to desired path if different
                if expected_pdf != pdf_path:
                    expected_pdf.rename(pdf_path)
                return True
            
            return False
            
        except (subprocess.TimeoutExpired, Exception) as e:
            print(f"PDF conversion error: {e}")
            return False


def create_sample_template():
    """Create a sample Arabic report template for testing"""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx not installed. Run: pip install python-docx")
        return None
    
    factory = DocumentFactory()
    template_path = factory.TEMPLATES_DIR / "eosb_report_template.docx"
    
    if template_path.exists():
        print(f"Template already exists: {template_path}")
        return template_path
    
    # Create new document
    doc = Document()
    
    # Title (Arabic - RTL)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("تقرير الخبير المحاسبي")
    run.bold = True
    run.font.size = Pt(18)
    
    # Case info
    doc.add_paragraph()
    doc.add_paragraph("الدعوى رقم: {{ case_number }}")
    doc.add_paragraph("المحكمة: {{ court_name }}")
    doc.add_paragraph("الولاية: {{ wilaya }}")
    
    # Parties
    doc.add_paragraph()
    doc.add_paragraph("المدعي: {{ plaintiff_name }}")
    doc.add_paragraph("المدعى عليه: {{ defendant_name }}")
    
    # Dates
    doc.add_paragraph()
    doc.add_paragraph("التاريخ: {{ gregorian_date }}م / {{ hijri_date }}هـ")
    
    # Calculation section
    doc.add_paragraph()
    calc_title = doc.add_paragraph()
    run = calc_title.add_run("حساب مستحقات نهاية الخدمة")
    run.bold = True
    
    doc.add_paragraph("مدة الخدمة: {{ service_period }}")
    doc.add_paragraph()
    
    # Results
    doc.add_paragraph("مكافأة نهاية الخدمة: {{ eosb }} ريال عماني")
    doc.add_paragraph("بدل الإشعار: {{ notice_pay }} ريال عماني")
    doc.add_paragraph("رصيد الإجازات: {{ leave_pay }} ريال عماني")
    
    # Total
    doc.add_paragraph()
    total_para = doc.add_paragraph()
    run = total_para.add_run("إجمالي المستحقات: {{ total_due }} ريال عماني")
    run.bold = True
    
    # Signature
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("الخبير: {{ expert_name }}")
    doc.add_paragraph("{{ expert_title }}")
    
    # Save template
    doc.save(template_path)
    print(f"Created template: {template_path}")
    return template_path


if __name__ == "__main__":
    from decimal import Decimal
    
    # Create sample template
    print("Creating sample template...")
    create_sample_template()
    
    # Test generation
    factory = DocumentFactory()
    print(f"\nTemplates available: {factory.list_templates()}")
    print(f"LibreOffice: {'Found' if factory.libreoffice_path else 'Not found'}")
    
    if factory.list_templates():
        # Test with sample data
        test_data = ReportData(
            case_number="1409/2024",
            court_name="المحكمة الابتدائية بمسقط",
            wilaya="مسقط",
            plaintiff_name="محمد بن عبدالله العماني",
            defendant_name="شركة عمان للتطوير",
            gregorian_date="2024-10-15",
            hijri_date="1446-04-12",
            assignment_date="2024-09-01",
            basic_salary=Decimal("750.000"),
            gross_salary=Decimal("1150.000"),
            eosb=Decimal("3633.904"),
            notice_pay=Decimal("750.000"),
            leave_pay=Decimal("775.000"),
            total_due=Decimal("5158.904"),
            years_of_service=6,
            months_of_service=4
        )
        
        result = factory.generate_report("eosb_report_template.docx", test_data)
        print(f"\nGeneration result: {result}")
